#!/usr/bin/env python
"""Convert a paper PDF to DOCX, keeping the figures and the reading order.

Why this is the preferred first step: a PDF-to-DOCX converter has already
solved the two hardest problems of this whole skill - the figures come across
as embedded images, and they land in roughly the right place in the text. No
figure extraction, no panel splitting, no relocating captions. You then
translate the body paragraphs in place and the layout survives.

Acrobat Pro is the best converter available, and this script drives it end to
end with no clicking. Getting there took three separate fixes, because a
single misleading COM error ("not implemented") hid all of them:

  1. pywin32 invokes IDispatch methods with DISPATCH_METHOD|DISPATCH_PROPERTYGET.
     Acrobat's JSObject answers that combination with E_NOTIMPL for *every*
     method it owns - privileged or not, getPageNumWords() included. A bare
     DISPATCH_METHOD works. See call(). This, not the privilege model, is what
     made the whole route look impossible.
  2. doc.saveAs really is privileged, so it has to come from a trusted
     function in a folder-level script. Acrobat 25.x ignores the per-user
     JavaScripts folder; only the application-level one next to debugger.js is
     read, and writing there needs elevation - hence the UAC prompt in
     install_acrobat_js(). Acrobat itself runs unelevated as usual.
  3. While Protected Mode is on, saveAs neither returns nor raises: it hangs
     forever. protected_mode_off() switches it off around the export only, and
     puts it back afterwards.

Tested on Acrobat Pro 25.1 (Exchange-Pro), pywin32 311, Windows 11.

Layout mode is the opposite of what you would guess. "Retain Page Layout"
(iLayoutMode=1) pins every block to its visual position, which shreds running
text into hundreds of text boxes, writes each twice (DrawingML + VML) and
reorders sentences across block boundaries - on a real paper it produced
"weakly allowed due to|transitions22,23. Notably,|orbital angular momentum
mixing". "Retain Flowing Text" (iLayoutMode=0) keeps reading order, headings
and paragraph breaks, and still places figures inline, which is what
docx_extract.py wants. Flowing is the default; --layout page is there if you
ever need visual fidelity instead.

RapidOCR is the automatic fallback. It rebuilds editable text and embeds
source-page previews for review; it does not claim to reproduce the layout.
The conversion report requires user review for this route, not for a
successful Acrobat export. These are the only two conversion engines.

Usage:
    python pdf_to_docx.py <pdf> [-o out.docx] [--engine acrobat|rapidocr|auto]
                                [--layout flowing|page]
    python pdf_to_docx.py --install-acrobat-js   # one time, prompts for UAC
    python pdf_to_docx.py --check

Exit codes:
    0  DOCX and .conversion.json written
    2  no converter available
    1  error
"""

import argparse
import contextlib
import csv
import ctypes
import io
import json
import multiprocessing
import os
import subprocess
import sys
import tempfile
import time
import zipfile
from queue import Empty

try:
    import winreg
except ImportError:
    winreg = None

try:
    import pythoncom
    import win32com.client as win32
except ImportError:                                 # reported by check()
    pythoncom = win32 = None

DC = r"SOFTWARE\Adobe\Adobe Acrobat\DC"
PRIV_KEY = DC + r"\Privileged"
DOCX_SETTINGS = DC + r"\AVConversionFromPDF\cSettings\c1\cSettings"

JS_NAME = "paper_translator.js"
JS_VERSION = "paper-translator/4"
PM_STASH = os.path.join(tempfile.gettempdir(), "paper_translator_pm_restore")

ACROBAT_JS = '''// paper-translator skill: trusted PDF -> DOCX export.
//
// Acrobat grants doc.saveAs only to a trusted function declared in a
// folder-level script. Acrobat 25.x ignores the per-user JavaScripts folder,
// so this file belongs in the APPLICATION folder, next to debugger.js:
//   <Acrobat install dir>\\Javascripts\\
// Writing there needs elevation; pdf_to_docx.py --install-acrobat-js does it.
//
// Errors come back as a return string, not a throw: a JS exception surfaces
// over COM as an opaque "server threw an exception" with no detail at all.

// Non-privileged probe. If COM can call this, the folder script is loaded.
var tpVersion = function () { return "%s"; };

// Export the document the COM caller already has open. Reopening that same
// PDF from inside Acrobat deadlocks, so this must not call app.openDoc.
var tpExportThis = app.trustedFunction(function (outPath, convID) {
    try {
        app.beginPriv();
        this.saveAs({ cPath: outPath,
                      cConvID: convID || "com.adobe.acrobat.docx" });
        app.endPriv();
        return "ok";
    } catch (e) {
        try { app.endPriv(); } catch (ePriv) {}
        return "ERROR: " + e;
    }
});
''' % JS_VERSION

NO_CONVERTER = """
No selected converter succeeded. Do not skip the capability gate.
If Acrobat Pro and RapidOCR are unusable, obtain the user's explicit consent
to download/install RapidOCR, repair the reported dependencies, and rerun
preflight.py. If consent is refused or pending, stop.
"""


# --- registry ---------------------------------------------------------------

def reg_get(sub, name):
    if winreg is None:
        return None
    try:
        with winreg.OpenKey(winreg.HKEY_CURRENT_USER, sub) as k:
            return winreg.QueryValueEx(k, name)[0]
    except OSError:
        return None


def reg_set(sub, name, value):
    with winreg.CreateKeyEx(winreg.HKEY_CURRENT_USER, sub, 0,
                            winreg.KEY_SET_VALUE) as k:
        winreg.SetValueEx(k, name, 0, winreg.REG_DWORD, value)


def restore_stashed_protected_mode():
    """Put Protected Mode back if an earlier run died mid-export."""
    if not os.path.exists(PM_STASH):
        return
    try:
        with open(PM_STASH) as f:
            value = int(f.read().strip())
        reg_set(PRIV_KEY, "bProtectedMode", value)
        print(f"note    : restored Protected Mode={value}, left off by an "
              f"earlier run")
    except (OSError, ValueError):
        pass
    with contextlib.suppress(OSError):
        os.remove(PM_STASH)


@contextlib.contextmanager
def protected_mode_off():
    """saveAs hangs forever with Protected Mode on - no error, no timeout.

    The original value is stashed on disk as well as in memory, so a killed
    process does not leave the sandbox switched off for good.
    """
    before = reg_get(PRIV_KEY, "bProtectedMode")
    if before in (None, 0):
        yield
        return
    with open(PM_STASH, "w") as f:
        f.write(str(before))
    reg_set(PRIV_KEY, "bProtectedMode", 0)
    try:
        yield
    finally:
        reg_set(PRIV_KEY, "bProtectedMode", before)
        with contextlib.suppress(OSError):
            os.remove(PM_STASH)


# --- the trusted script -----------------------------------------------------

def acrobat_exe():
    """Acrobat's install path, from its App Paths registry entry."""
    if winreg is None:
        return None
    key = (r"SOFTWARE\Microsoft\Windows\CurrentVersion\App Paths"
           r"\Acrobat.exe")
    for hive in (winreg.HKEY_LOCAL_MACHINE, winreg.HKEY_CURRENT_USER):
        try:
            with winreg.OpenKey(hive, key) as k:
                path = winreg.QueryValueEx(k, "")[0]
        except OSError:
            continue
        if os.path.isfile(path):
            return path
    return None


def app_js_path():
    """The application-level folder script Acrobat 25.x actually reads."""
    exe = acrobat_exe()
    if not exe:
        return None
    return os.path.join(os.path.dirname(exe), "Javascripts", JS_NAME)


def js_current(path):
    """True if the installed script is byte-identical to what we ship."""
    try:
        with open(path, encoding="utf-8") as f:
            return f.read() == ACROBAT_JS
    except OSError:
        return False


def install_acrobat_js(quiet=False):
    """Copy the trusted-function script into Acrobat's install folder.

    That folder is under Program Files, so this needs elevation and raises a
    UAC prompt. Only the copy is elevated - Acrobat and this script keep
    running as the normal user, which is what COM needs.
    """
    dst = app_js_path()
    if not dst:
        print("ERROR: cannot locate Acrobat's install folder (Acrobat Pro "
              "installed?)", file=sys.stderr)
        return False
    if js_current(dst):
        if not quiet:
            print(f"already current: {dst}")
        return True

    src = os.path.join(tempfile.gettempdir(), JS_NAME)
    with open(src, "w", encoding="utf-8") as f:
        f.write(ACROBAT_JS)
    # Driving the copy from a .ps1 keeps these paths out of a nested-quoted
    # command line. Single quotes: PowerShell treats backslashes literally.
    ps1 = os.path.join(tempfile.gettempdir(), "paper_translator_install.ps1")
    with open(ps1, "w", encoding="utf-8") as f:
        f.write("$ErrorActionPreference = 'Stop'\n")
        f.write("New-Item -ItemType Directory -Force -Path '%s' | Out-Null\n"
                % os.path.dirname(dst))
        f.write("Copy-Item -LiteralPath '%s' -Destination '%s' -Force\n"
                % (src, dst))

    print("elevating: approve the UAC prompt so Acrobat can load the export")
    print("           script (one time; Acrobat itself stays unelevated)")
    rc = ctypes.windll.shell32.ShellExecuteW(
        None, "runas", "pwsh.exe",
        '-NoProfile -ExecutionPolicy Bypass -File "%s"' % ps1, None, 0)
    if rc <= 32:
        print(f"ERROR: elevation refused or failed (ShellExecute {rc})",
              file=sys.stderr)
        return False
    for _ in range(90):             # ShellExecuteW returns before the copy
        if js_current(dst):
            print(f"installed: {dst}")
            return True
        time.sleep(1)
    print("ERROR: the elevated copy never completed", file=sys.stderr)
    return False


# --- COM plumbing -----------------------------------------------------------

def call(obj, name, *args):
    """Invoke a JSObject method with a bare DISPATCH_METHOD.

    pywin32's normal `obj.name(...)` adds DISPATCH_PROPERTYGET to the flags,
    and Acrobat answers that with E_NOTIMPL for every method it has. Reading
    properties works either way; only calls need this.
    """
    dispid = obj._oleobj_.GetIDsOfNames(0, name)
    return obj._oleobj_.Invoke(dispid, 0, pythoncom.DISPATCH_METHOD, True,
                               *args)


def devpath(win_path):
    """C:\\a\\b.pdf -> /C/a/b.pdf, the only form Acrobat JS accepts."""
    p = os.path.abspath(win_path)
    return "/" + p[0] + p[2:].replace("\\", "/")


def kill_acrobat():
    for exe in ("Acrobat.exe", "AcroCEF.exe"):
        subprocess.run(["taskkill", "/IM", exe, "/F"], capture_output=True)
    time.sleep(3)


def _acrobat_pids():
    """Snapshot Acrobat processes so timeout cleanup leaves older ones alone."""
    result = subprocess.run(
        ["tasklist", "/FI", "IMAGENAME eq Acrobat.exe", "/FO", "CSV", "/NH"],
        capture_output=True, text=True, errors="replace", check=True, timeout=10)
    return {int(row[1]) for row in csv.reader(result.stdout.splitlines())
            if len(row) >= 2 and row[0].lower() == "acrobat.exe"}


def _acrobat_worker(pdf, out, layout, result):
    result.put(_export_acrobat(pdf, out, layout))


def convert_acrobat(pdf, out, layout="flowing", install=True, restart=True,
                    timeout=180):
    """Export through Acrobat with a bounded, isolated COM worker."""
    if win32 is None:
        print("  acrobat: pywin32 missing (pip install pywin32)")
        return False
    dst = app_js_path()
    if not dst:
        print("  acrobat: not installed")
        return False
    if not js_current(dst):
        if not install:
            print("  acrobat: export script not installed "
                  "(run --install-acrobat-js)")
            return False
        if not install_acrobat_js(quiet=True):
            return False

    reg_set(DOCX_SETTINGS, "iLayoutMode", 1 if layout == "page" else 0)
    reg_set(DOCX_SETTINGS, "bIncludeImages", 1)

    # Protected Mode is read at startup, so the restart has to happen inside
    # the context manager, not before it.
    with protected_mode_off():
        if restart:
            kill_acrobat()
        before = _acrobat_pids()
        context = multiprocessing.get_context("spawn")
        result = context.Queue()
        worker = context.Process(
            target=_acrobat_worker, args=(pdf, out, layout, result), daemon=True)
        try:
            worker.start()
            worker.join(timeout)
            if worker.is_alive():
                print(f"  acrobat: export timed out after {timeout:g}s; "
                      "stopping this attempt")
                return False
            try:
                return bool(result.get(timeout=1))
            except Empty:
                print(f"  acrobat: worker exited without a result "
                      f"(exit code {worker.exitcode})")
                return False
        finally:
            if worker.is_alive():
                worker.terminate()
                worker.join(5)
                for pid in _acrobat_pids() - before:
                    subprocess.run(
                        ["taskkill", "/PID", str(pid), "/T", "/F"],
                        capture_output=True)
            result.close()


def _export_acrobat(pdf, out, layout):
    """COM calls may block on Acrobat UI; the parent owns the timeout."""
    app = avdoc = None
    try:
        app = win32.DispatchEx("AcroExch.App")
        avdoc = win32.DispatchEx("AcroExch.AVDoc")
        if not avdoc.Open(os.path.abspath(pdf), "paper-translator"):
            print("  acrobat: could not open the PDF", flush=True)
            return False
        jso = avdoc.GetPDDoc().GetJSObject()
        try:
            version = call(jso, "tpVersion")
        except Exception:                       # noqa: BLE001
            print("  acrobat: folder script not loaded - quit Acrobat "
                  "completely and rerun --install-acrobat-js", flush=True)
            return False
        print(f"  acrobat: script {version}, layout {layout}", flush=True)
        result = call(jso, "tpExportThis", devpath(out))
        if result != "ok":
            print(f"  acrobat: {result}", flush=True)
            return False
        for _ in range(30):
            if os.path.exists(out) and os.path.getsize(out) > 0:
                return True
            time.sleep(1)
        print("  acrobat: saveAs reported ok but wrote no file", flush=True)
        return False
    except Exception as exc:
        print(f"  acrobat: {type(exc).__name__}: {exc}", flush=True)
        return False
    finally:
        with contextlib.suppress(Exception):
            avdoc.Close(True)
        with contextlib.suppress(Exception):
            app.Exit()


def convert_rapidocr(pdf, out, lang="en", dpi=200):
    """Rebuild editable OCR text with embedded source pages for user review."""
    try:
        import pymupdf as fitz
        import numpy as np
        from docx import Document
        from docx.shared import Inches, Pt
        import ocr_engine
    except ImportError as exc:
        print(f"  rapidocr: missing dependency: {exc}")
        return False

    engine, _, note = ocr_engine.make_engine(
        lang=lang, prefer=ocr_engine.BACKEND_RAPID)
    if engine is None:
        print(f"  rapidocr: {note}")
        return False

    doc = Document()
    doc.core_properties.subject = "paper-translator:rapidocr; user review required"
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(11)
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.75)
    width = section.page_width - section.left_margin - section.right_margin
    height = section.page_height - section.top_margin - section.bottom_margin
    total_chars = 0
    try:
        with fitz.open(pdf) as source:
            if not len(source) or source.needs_pass:
                raise ValueError("PDF is empty or requires a password")
            for index, page in enumerate(source):
                pix = page.get_pixmap(dpi=dpi, colorspace=fitz.csRGB, alpha=False)
                rgb = np.frombuffer(pix.samples, dtype=np.uint8).reshape(
                    pix.height, pix.width, 3)
                rows = ocr_engine.read(engine, rgb)
                lines = [row["text"].strip() for row in rows if row["text"].strip()]
                if not lines:
                    raise ValueError(
                        f"page {index + 1}: no text recognized; review the scan "
                        "or rerun at a higher --dpi")
                if index:
                    doc.add_page_break()
                doc.add_heading(f"Page {index + 1} - OCR text", level=1)
                for line in lines:
                    doc.add_paragraph(line)
                total_chars += sum(map(len, lines))
                doc.add_page_break()
                doc.add_heading(f"Page {index + 1} - source preview", level=1)
                preview_width = min(width, int((height - Inches(0.7))
                                              * pix.width / pix.height))
                doc.add_picture(io.BytesIO(pix.tobytes("png")),
                                width=preview_width)
                print(f"  rapidocr: page {index + 1}/{len(source)}, "
                      f"{sum(map(len, lines))} chars", flush=True)
        doc.save(out)
    except Exception as exc:
        print(f"  rapidocr: {type(exc).__name__}: {exc}")
        return False
    print(f"  rapidocr: {total_chars} editable chars; source previews embedded")
    return True


def valid_docx(path):
    """Reject missing/truncated exports instead of reporting false success."""
    try:
        with zipfile.ZipFile(path) as archive:
            return ("word/document.xml" in archive.namelist()
                    and archive.testzip() is None)
    except (OSError, zipfile.BadZipFile):
        return False


def check():
    """Check Acrobat and RapidOCR without launching Word or another app."""
    import preflight

    print("converter availability")
    probes = preflight.collect()
    for probe in preflight.figure_caps(probes):
        preflight._row(probe)

    exe = acrobat_exe()
    if exe:
        print(f"  Acrobat Pro      : {exe}")
        dst = app_js_path()
        state = ("current" if js_current(dst) else
                 "STALE - will reinstall" if os.path.exists(dst) else
                 "not installed - will prompt for UAC")
        print(f"  export script    : {state}")
        print(f"                     {dst}")
    else:
        print("  Acrobat Pro      : no")

    pm = reg_get(PRIV_KEY, "bProtectedMode")
    print(f"  Protected Mode   : {pm} "
          f"{'(switched off during export, then restored)' if pm else ''}")
    layout = reg_get(DOCX_SETTINGS, "iLayoutMode")
    print(f"  iLayoutMode      : {layout} "
          f"({'page' if layout == 1 else 'flowing'}); set per export")

    return preflight.decide(probes)


def main():
    with contextlib.suppress(Exception):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser(add_help=True)
    ap.add_argument("pdf", nargs="?")
    ap.add_argument("-o", "--out")
    ap.add_argument("--engine", choices=("auto", "acrobat", "rapidocr"),
                    default="auto")
    ap.add_argument("--ocr-lang", default="en", help="RapidOCR language")
    ap.add_argument("--dpi", type=int, default=200, help="RapidOCR render DPI")
    ap.add_argument("--acrobat-timeout", type=float, default=180,
                    help="maximum Acrobat export seconds before fallback (default: 180)")
    ap.add_argument("--layout", choices=("flowing", "page"), default="flowing",
                    help="flowing keeps reading order (default); page keeps "
                         "visual position but scrambles sentence order")
    ap.add_argument("--install-acrobat-js", action="store_true",
                    help="install the trusted export script (prompts for UAC)")
    ap.add_argument("--no-install", action="store_true",
                    help="fail instead of prompting for UAC when the export "
                         "script is missing")
    ap.add_argument("--check", action="store_true",
                    help="report which converters are available")
    ap.add_argument("--no-restart", action="store_true",
                    help="do not kill a running Acrobat before exporting")
    args = ap.parse_args()
    if args.dpi <= 0:
        ap.error("--dpi must be positive")
    if args.acrobat_timeout <= 0:
        ap.error("--acrobat-timeout must be positive")

    if args.install_acrobat_js:
        return 0 if install_acrobat_js() else 1
    if args.check:
        return check()
    if not args.pdf:
        ap.error("give a PDF, or use --check / --install-acrobat-js")
    if not os.path.isfile(args.pdf):
        print(f"ERROR: no such file: {args.pdf}", file=sys.stderr)
        return 1

    out = args.out or os.path.splitext(os.path.abspath(args.pdf))[0] + ".docx"
    report_path = os.path.splitext(os.path.abspath(out))[0] + ".conversion.json"
    if os.path.exists(out) or os.path.exists(report_path):
        print(f"ERROR: output already exists; choose an unused -o path: {out}",
              file=sys.stderr)
        return 1
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    order = {"auto": ("acrobat", "rapidocr"),
             "acrobat": ("acrobat",), "rapidocr": ("rapidocr",)}[args.engine]

    print(f"pdf     : {args.pdf}")
    for engine in order:
        print(f"trying  : {engine}")
        if engine == "acrobat":
            restore_stashed_protected_mode()
            ok = convert_acrobat(args.pdf, out, layout=args.layout,
                                 install=not args.no_install,
                                 restart=not args.no_restart,
                                 timeout=args.acrobat_timeout)
        else:
            ok = convert_rapidocr(args.pdf, out, lang=args.ocr_lang, dpi=args.dpi)
        if ok and valid_docx(out):
            review = engine == "rapidocr"
            with open(report_path, "w", encoding="utf-8") as stream:
                json.dump({
                    "pdf": os.path.abspath(args.pdf),
                    "docx": os.path.abspath(out),
                    "engine": engine,
                    "requires_word_review": review,
                }, stream, ensure_ascii=False, indent=2)
            print(f"docx    : {out}  "
                  f"({os.path.getsize(out) / 1e6:.1f} MB, via {engine})")
            print(f"report  : {report_path}")
            print("review  : " + (
                "REQUIRED - show the Word result and ask whether it is "
                "satisfactory; do not continue before the user's answer"
                if review else
                "SKIP - Acrobat export succeeded; do not ask Word satisfaction"))
            return 0
        if ok:
            print(f"  {engine}: invalid DOCX output")
    print("\nno converter worked.")
    print(NO_CONVERTER)
    return 2


if __name__ == "__main__":
    sys.exit(main())
