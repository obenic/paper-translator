import contextlib
import io
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import pdf_to_docx


class ConversionTests(unittest.TestCase):
    def run_conversion(self, acrobat_ok):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "source.pdf"
            source.write_bytes(b"%PDF-test")
            target = Path(tmp) / "result.docx"

            def convert(_pdf, out, **_kwargs):
                doc = Document()
                doc.add_paragraph("Verified conversion fixture.")
                doc.save(out)
                return True

            with patch.object(sys, "argv", [
                    "pdf_to_docx.py", str(source), "-o", str(target),
                    "--no-install", "--no-restart",
            ]), patch.object(pdf_to_docx, "restore_stashed_protected_mode"), \
                    patch.object(pdf_to_docx, "convert_acrobat",
                                 side_effect=convert if acrobat_ok else None,
                                 return_value=False), \
                    patch.object(pdf_to_docx, "convert_rapidocr",
                                 side_effect=convert, create=True), \
                    contextlib.redirect_stdout(io.StringIO()):
                code = pdf_to_docx.main()

            self.assertEqual(code, 0)
            report = target.with_suffix(".conversion.json")
            self.assertTrue(report.exists(), "successful conversion must record its source")
            return json.loads(report.read_text(encoding="utf-8"))

    def test_acrobat_success_skips_word_review(self):
        report = self.run_conversion(acrobat_ok=True)
        self.assertEqual(report["engine"], "acrobat")
        self.assertIs(report["requires_word_review"], False)

    def test_acrobat_failure_uses_rapidocr_and_requires_review(self):
        report = self.run_conversion(acrobat_ok=False)
        self.assertEqual(report["engine"], "rapidocr")
        self.assertIs(report["requires_word_review"], True)

    def test_explicit_rapidocr_does_not_touch_acrobat_protection(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "source.pdf"
            source.write_bytes(b"%PDF-test")
            target = Path(tmp) / "result.docx"

            def convert(_pdf, out, **_kwargs):
                doc = Document()
                doc.add_paragraph("OCR result.")
                doc.save(out)
                return True

            with patch.object(sys, "argv", [
                    "pdf_to_docx.py", str(source), "-o", str(target),
                    "--engine", "rapidocr",
            ]), patch.object(
                    pdf_to_docx, "restore_stashed_protected_mode",
                    side_effect=AssertionError("OCR must not touch Acrobat state"),
            ), patch.object(pdf_to_docx, "convert_rapidocr", side_effect=convert), \
                    contextlib.redirect_stdout(io.StringIO()):
                self.assertEqual(pdf_to_docx.main(), 0)

    def test_word_engine_is_rejected(self):
        errors = io.StringIO()
        with patch.object(sys, "argv", [
                "pdf_to_docx.py", "--engine", "word",
        ]), contextlib.redirect_stderr(errors):
            with self.assertRaises(SystemExit) as result:
                pdf_to_docx.main()
        self.assertEqual(result.exception.code, 2)
        self.assertIn("invalid choice", errors.getvalue())
        self.assertFalse(hasattr(pdf_to_docx, "convert_word"))

    def test_failed_converters_do_not_report_success(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "source.pdf"
            source.write_bytes(b"%PDF-test")
            target = Path(tmp) / "result.docx"
            with patch.object(sys, "argv", [
                    "pdf_to_docx.py", str(source), "-o", str(target),
            ]), patch.object(pdf_to_docx, "restore_stashed_protected_mode"), \
                    patch.object(pdf_to_docx, "convert_acrobat", return_value=True), \
                    patch.object(pdf_to_docx, "convert_rapidocr", return_value=False), \
                    contextlib.redirect_stdout(io.StringIO()):
                self.assertEqual(pdf_to_docx.main(), 2)
            self.assertFalse(target.with_suffix(".conversion.json").exists())

    def test_acrobat_timeout_restores_protection_and_keeps_existing_processes(self):
        context = Mock()
        worker = context.Process.return_value
        worker.is_alive.return_value = True
        with tempfile.TemporaryDirectory() as tmp, \
                patch.object(pdf_to_docx, "PM_STASH", str(Path(tmp) / "restore")), \
                patch.object(pdf_to_docx, "win32", Mock()), \
                patch.object(pdf_to_docx, "app_js_path", return_value="export.js"), \
                patch.object(pdf_to_docx, "js_current", return_value=True), \
                patch.object(pdf_to_docx, "reg_get", return_value=1), \
                patch.object(pdf_to_docx, "reg_set") as reg_set, \
                patch.object(pdf_to_docx, "_acrobat_pids",
                             side_effect=[{11}, {11, 22}], create=True), \
                patch("multiprocessing.get_context", return_value=context), \
                patch.object(pdf_to_docx.subprocess, "run") as run, \
                contextlib.redirect_stdout(io.StringIO()):
            self.assertFalse(pdf_to_docx.convert_acrobat(
                "source.pdf", "target.docx", restart=False, timeout=0.01))
            self.assertFalse(Path(pdf_to_docx.PM_STASH).exists())
        worker.terminate.assert_called_once()
        worker.join.assert_any_call(0.01)
        run.assert_called_once_with(
            ["taskkill", "/PID", "22", "/T", "/F"], capture_output=True)
        self.assertEqual(reg_set.call_args.args,
                         (pdf_to_docx.PRIV_KEY, "bProtectedMode", 1))


if __name__ == "__main__":
    unittest.main()
